#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
中西医映射关系字典构建器
基于SymMap数据构建Syndrome到Disease的映射关系
"""

import json
import os
from pathlib import Path
from openpyxl import load_workbook


class MappingDictBuilder:
    """中西医映射关系字典构建器"""
    
    def __init__(self, resource_dir: str):
        """
        初始化构建器
        
        Args:
            resource_dir: SymMap资源文件目录
        """
        self.resource_dir = Path(resource_dir)
        self.mapping_dict = {}
        
    def read_excel_file(self, filename: str, max_rows: int = None):
        """
        读取Excel文件
        
        Args:
            filename: 文件名
            max_rows: 最大读取行数（None表示全部）
            
        Returns:
            (headers, data): 列名列表和数据列表
        """
        file_path = self.resource_dir / filename
        
        if not file_path.exists():
            print(f"❌ 文件不存在: {file_path}")
            return None, None
        
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            
            # 读取表头
            headers = [cell.value for cell in ws[1]]
            
            # 读取数据
            data = []
            for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if max_rows and i > max_rows + 1:
                    break
                data.append(row)
            
            wb.close()
            return headers, data
            
        except Exception as e:
            print(f"❌ 读取文件失败: {e}")
            return None, None
    
    def build_syndrome_to_disease_mapping(self):
        """构建证候到疾病的映射关系"""
        print("\n=== 开始构建证候到疾病映射关系 ===\n")
        
        # 读取数据文件
        print("正在读取证候群数据...")
        headers, data = self.read_excel_file("SymMap v2.0, SMSY （证候群）file.xlsx")
        
        if headers is None or data is None:
            print("❌ 无法读取证候群文件")
            return
        
        print(f"✅ 成功读取证候群文件，共 {len(data)} 条记录")
        print(f"列名: {headers}")
        
        # 构建映射字典
        syndrome_to_disease = {}
        
        # TODO: 根据实际数据结构构建映射
        # 这里需要先查看数据结构
        
        self.mapping_dict['Syndrome_to_Disease'] = syndrome_to_disease
        print(f"\n✅ 成功构建 {len(syndrome_to_disease)} 个证候的映射关系")
    
    def build_full_mapping_dict(self):
        """构建完整的映射字典"""
        print("=== 开始构建完整映射关系字典 ===\n")
        
        # 1. 证候到疾病映射
        self.build_syndrome_to_disease_mapping()
        
        # 2. 其他映射关系可以继续添加
        # TODO: 添加其他映射关系
        
    def save_to_json(self, output_file: str):
        """
        保存映射字典到JSON文件
        
        Args:
            output_file: 输出文件路径
        """
        print(f"\n正在保存映射字典到: {output_file}")
        
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(self.mapping_dict, f, ensure_ascii=False, indent=2)
        
        print(f"✅ 映射字典已保存")
    
    def display_sample_data(self):
        """显示所有Excel文件的列名和示例数据"""
        print("\n=== 分析SymMap数据结构 ===\n")
        
        excel_files = [
            ("证候群", "SymMap v2.0, SMSY （证候群）file.xlsx"),
            ("疾病", "SymMap v2.0, SMDE (疾病)file.xlsx"),
            ("中医症状", "SymMap v2.0, SMTS（中医症状） file.xlsx"),
            ("现代医学症状", "SymMap v2.0, SMMS (MM症状)file.xlsx"),
            ("药草", "SymMap v2.0, SMHB（药草） file.xlsx"),
            ("成分", "SymMap v2.0, SMIT（成分） file.xlsx"),
            ("靶点", "SymMap v2.0, SMTT（目标） file.xlsx"),
        ]
        
        for name, filename in excel_files:
            print(f"\n{'='*60}")
            print(f"📁 {name} ({filename})")
            print(f"{'='*60}")
            
            headers, data = self.read_excel_file(filename, max_rows=5)
            
            if headers is None:
                print(f"⚠️  文件不存在或读取失败")
                continue
            
            print(f"列名 ({len(headers)}列): {headers}")
            print(f"\n前5行数据:")
            
            for i, row in enumerate(data[:5], start=1):
                print(f"\n第{i}行:")
                for j, (col_name, value) in enumerate(zip(headers, row)):
                    if value is not None and value != '':
                        print(f"  {col_name}: {value}")
            print()


def main():
    """主函数"""
    # 切换到脚本所在目录
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    
    # 资源文件目录
    resource_dir = "./resoure"
    
    # 创建映射字典构建器
    builder = MappingDictBuilder(resource_dir)
    
    # 首先分析数据结构
    print("=" * 80)
    print("第一步：分析SymMap数据结构")
    print("=" * 80)
    builder.display_sample_data()
    
    print("\n" + "=" * 80)
    print("提示：请查看上面的数据结构，然后我将根据实际列名构建映射关系")
    print("=" * 80)


if __name__ == "__main__":
    main()
