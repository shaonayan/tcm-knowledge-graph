#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
中西医映射关系字典构建器
基于SymMap数据和现有的知识图谱三元组构建完整的映射关系
"""

import json
import os
from pathlib import Path
from openpyxl import load_workbook
from collections import defaultdict


class TCMWMMapper:
    """中西医映射关系构建器"""
    
    def __init__(self, resource_dir: str, triples_dir: str):
        """
        初始化构建器
        
        Args:
            resource_dir: SymMap资源文件目录
            triples_dir: 三元组数据目录
        """
        self.resource_dir = Path(resource_dir)
        self.triples_dir = Path(triples_dir)
        
        # 映射字典
        self.syndrome_to_disease = defaultdict(list)
        self.disease_to_symptom = defaultdict(list)
        self.syndrome_to_symptom = defaultdict(list)
        
    def load_symmap_syndrome_data(self):
        """加载SymMap证候数据"""
        print("正在加载SymMap证候数据...")
        
        file_path = self.resource_dir / "SymMap v2.0, SMSY （证候群）file.xlsx"
        
        if not file_path.exists():
            print(f"❌ 文件不存在: {file_path}")
            return {}
        
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            
            # 读取表头
            headers = [cell.value for cell in ws[1]]
            
            # 构建证候字典
            syndrome_dict = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = dict(zip(headers, row))
                syndrome_id = row_dict.get('Syndrome_id')
                syndrome_name = row_dict.get('Syndrome_name')
                
                if syndrome_id and syndrome_name:
                    syndrome_dict[syndrome_name] = {
                        'id': syndrome_id,
                        'name': syndrome_name,
                        'english': row_dict.get('Syndrome_English', ''),
                        'pinyin': row_dict.get('Syndrome_PinYin', ''),
                        'definition': row_dict.get('Syndrome_definition', '')
                    }
            
            wb.close()
            print(f"✅ 成功加载 {len(syndrome_dict)} 个证候")
            return syndrome_dict
            
        except Exception as e:
            print(f"❌ 加载证候数据失败: {e}")
            return {}
    
    def load_symmap_disease_data(self):
        """加载SymMap疾病数据"""
        print("正在加载SymMap疾病数据...")
        
        file_path = self.resource_dir / "SymMap v2.0, SMDE (疾病)file.xlsx"
        
        if not file_path.exists():
            print(f"❌ 文件不存在: {file_path}")
            return {}
        
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            
            # 读取表头
            headers = [cell.value for cell in ws[1]]
            
            # 构建疾病字典
            disease_dict = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = dict(zip(headers, row))
                disease_id = row_dict.get('Disease_id')
                disease_name = row_dict.get('Disease_name')
                
                if disease_id and disease_name:
                    disease_dict[disease_name] = {
                        'id': disease_id,
                        'name': disease_name,
                        'english': row_dict.get('Disease_English', ''),
                        'category': row_dict.get('Category', '')
                    }
            
            wb.close()
            print(f"✅ 成功加载 {len(disease_dict)} 个疾病")
            return disease_dict
            
        except Exception as e:
            print(f"❌ 加载疾病数据失败: {e}")
            return {}
    
    def load_triples_data(self):
        """加载已有的三元组数据"""
        print("\n正在加载三元组数据...")
        
        # 加载中医三元组
        chinese_file = self.triples_dir / "chinese_triples.json"
        western_file = self.triples_dir / "western_triples.json"
        
        chinese_triples = []
        western_triples = []
        
        if chinese_file.exists():
            with open(chinese_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                chinese_triples = data.get('triples', [])
            print(f"✅ 加载中医三元组: {len(chinese_triples)} 条")
        
        if western_file.exists():
            with open(western_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                western_triples = data.get('triples', [])
            print(f"✅ 加载西医三元组: {len(western_triples)} 条")
        
        return chinese_triples + western_triples
    
    def build_syndrome_to_disease_mapping(self, triples):
        """从三元组构建证候到疾病的映射"""
        print("\n构建证候到疾病映射...")
        
        count = 0
        for triple in triples:
            head = triple.get('head', '')
            relation = triple.get('relation', '')
            tail = triple.get('tail', '')
            
            # 找到"对应"关系（证候 -> 疾病）
            if relation == '对应':
                if tail not in self.syndrome_to_disease[head]:
                    self.syndrome_to_disease[head].append(tail)
                    count += 1
        
        print(f"✅ 构建了 {len(self.syndrome_to_disease)} 个证候的映射，共 {count} 条关系")
    
    def build_disease_to_symptom_mapping(self, triples):
        """从三元组构建疾病到症状的映射"""
        print("\n构建疾病到症状映射...")
        
        count = 0
        for triple in triples:
            head = triple.get('head', '')
            relation = triple.get('relation', '')
            tail = triple.get('tail', '')
            
            # 找到"表现为"关系（疾病 -> 症状）
            if relation == '表现为':
                if tail not in self.disease_to_symptom[head]:
                    self.disease_to_symptom[head].append(tail)
                    count += 1
        
        print(f"✅ 构建了 {len(self.disease_to_symptom)} 个疾病的映射，共 {count} 条关系")
    
    def build_syndrome_to_symptom_mapping(self, triples):
        """从三元组构建证候到症状的映射"""
        print("\n构建证候到症状映射...")
        
        count = 0
        for triple in triples:
            head = triple.get('head', '')
            relation = triple.get('relation', '')
            tail = triple.get('tail', '')
            
            # 找到"症候表现"关系（证候 -> 症状）
            if relation == '症候表现':
                if tail not in self.syndrome_to_symptom[head]:
                    self.syndrome_to_symptom[head].append(tail)
                    count += 1
        
        print(f"✅ 构建了 {len(self.syndrome_to_symptom)} 个证候的映射，共 {count} 条关系")
    
    def generate_mapping_dict(self):
        """生成完整的映射字典"""
        print("\n=== 开始构建中西医映射字典 ===\n")
        
        # 1. 加载SymMap数据（可选，用于扩展）
        syndrome_dict = self.load_symmap_syndrome_data()
        disease_dict = self.load_symmap_disease_data()
        
        # 2. 加载三元组数据
        triples = self.load_triples_data()
        
        # 3. 构建各种映射关系
        self.build_syndrome_to_disease_mapping(triples)
        self.build_disease_to_symptom_mapping(triples)
        self.build_syndrome_to_symptom_mapping(triples)
        
        # 4. 构建最终的映射字典
        mapping_dict = {
            "Syndrome_to_Disease": dict(self.syndrome_to_disease),
            "Disease_to_Symptom": dict(self.disease_to_symptom),
            "Syndrome_to_Symptom": dict(self.syndrome_to_symptom)
        }
        
        return mapping_dict
    
    def save_mapping_dict(self, mapping_dict: dict, output_file: str):
        """保存映射字典到JSON文件"""
        print(f"\n正在保存映射字典到: {output_file}")
        
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(mapping_dict, f, ensure_ascii=False, indent=2)
        
        # 统计信息
        print("\n映射字典统计:")
        for key, value in mapping_dict.items():
            total_mappings = sum(len(v) if isinstance(v, list) else 0 for v in value.values())
            print(f"  {key}: {len(value)} 个实体，共 {total_mappings} 条映射关系")
        
        print(f"\n✅ 映射字典已保存到: {output_file}")
    
    def display_sample_mappings(self, mapping_dict: dict, sample_size: int = 5):
        """显示示例映射"""
        print(f"\n=== 映射关系示例 (前{sample_size}个) ===\n")
        
        for mapping_type, mappings in mapping_dict.items():
            print(f"\n{mapping_type}:")
            print("-" * 60)
            
            count = 0
            for entity, related_entities in mappings.items():
                if count >= sample_size:
                    break
                if related_entities:  # 只显示有映射关系的
                    print(f'  "{entity}": {related_entities}')
                    count += 1


def main():
    """主函数"""
    # 切换到脚本所在目录
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    
    # 配置路径
    resource_dir = "./resoure"
    triples_dir = "../../Bert-Crf模型调用/输出结果"
    output_file = "./中西医映射关系字典.json"
    
    # 创建映射构建器
    mapper = TCMWMMapper(resource_dir, triples_dir)
    
    # 生成映射字典
    mapping_dict = mapper.generate_mapping_dict()
    
    # 显示示例
    mapper.display_sample_mappings(mapping_dict, sample_size=10)
    
    # 保存映射字典
    mapper.save_mapping_dict(mapping_dict, output_file)
    
    print("\n🎉 中西医映射关系字典构建完成！")


if __name__ == "__main__":
    main()
